"""导入情感词典，并在文本中执行可解释的最长词匹配。"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Mapping


EMOTION_MAJOR_MAP = {
    "PA": "乐", "PE": "乐",
    "PD": "好", "PH": "好", "PG": "好", "PB": "好", "PK": "好",
    "NA": "怒",
    "NB": "哀", "NJ": "哀", "NH": "哀", "PF": "哀",
    "NI": "惧", "NC": "惧", "NG": "惧",
    "NE": "恶", "ND": "恶", "NN": "恶", "NK": "恶", "NL": "恶",
    "PC": "惊",
}

HEADER_ALIASES = {
    "word": ("word", "词语", "词", "单词"),
    "pos": ("pos", "词性种类", "词性"),
    "emotion": ("emotion", "category", "情感分类", "情绪分类"),
    "intensity": ("intensity", "强度", "情感强度"),
    "polarity": ("polarity", "极性", "情感极性"),
}


def normalize_emotion(category: object) -> str:
    """把 21 个细分类代码归并为 7 个易读的大类。"""
    code = str(category or "").strip().upper()
    return EMOTION_MAJOR_MAP.get(code, code if code in "乐好怒哀惧恶惊" else "其他")


def normalize_polarity(value: object) -> str:
    raw = str(value or "").strip().lower()
    if raw in {"1", "positive", "正", "正面", "褒义"}:
        return "positive"
    if raw in {"2", "negative", "负", "负面", "贬义"}:
        return "negative"
    if raw in {"3", "both", "双极性", "兼有"}:
        return "both"
    return "neutral"


@dataclass(frozen=True, slots=True)
class EmotionEntry:
    word: str
    pos: str
    category: str
    major_emotion: str
    intensity: float
    polarity: str


@dataclass(frozen=True, slots=True)
class EmotionMatch:
    word: str
    start: int
    end: int
    emotion: str
    category: str
    intensity: float
    polarity: str


def _pick(row: Mapping[str, object], field: str, default: object = "") -> object:
    for name in HEADER_ALIASES[field]:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return default


def entry_from_row(row: Mapping[str, object]) -> EmotionEntry:
    word = str(_pick(row, "word")).strip()
    if not word:
        raise ValueError("词库记录中的词语不能为空")
    category = str(_pick(row, "emotion")).strip().upper()
    try:
        intensity = float(_pick(row, "intensity", 1))
    except (TypeError, ValueError) as error:
        raise ValueError(f"词语“{word}”的强度不是数字") from error
    return EmotionEntry(
        word=word,
        pos=str(_pick(row, "pos")).strip(),
        category=category,
        major_emotion=normalize_emotion(category),
        intensity=max(0.0, intensity),
        polarity=normalize_polarity(_pick(row, "polarity")),
    )


class EmotionLexicon:
    """按首字符建立索引，扫描文本时优先匹配更长的短语。"""

    def __init__(self, entries: Iterable[EmotionEntry]):
        unique: dict[str, EmotionEntry] = {}
        for entry in entries:
            # 同词多义时保留强度更高的条目，规则明确且结果稳定。
            if entry.word not in unique or entry.intensity > unique[entry.word].intensity:
                unique[entry.word] = entry
        self.entries = sorted(unique.values(), key=lambda item: item.word)
        self._index: dict[str, list[EmotionEntry]] = {}
        for entry in self.entries:
            self._index.setdefault(entry.word[0], []).append(entry)
        for candidates in self._index.values():
            candidates.sort(key=lambda item: (-len(item.word), item.word))

    @classmethod
    def from_rows(cls, rows: Iterable[Mapping[str, object]]) -> "EmotionLexicon":
        return cls(entry_from_row(row) for row in rows)

    def __len__(self) -> int:
        return len(self.entries)

    def match(self, text: str) -> list[EmotionMatch]:
        matches: list[EmotionMatch] = []
        cursor = 0
        while cursor < len(text):
            selected = next(
                (entry for entry in self._index.get(text[cursor], []) if text.startswith(entry.word, cursor)),
                None,
            )
            if selected is None:
                cursor += 1
                continue
            end = cursor + len(selected.word)
            matches.append(EmotionMatch(
                word=selected.word,
                start=cursor,
                end=end,
                emotion=selected.major_emotion,
                category=selected.category,
                intensity=selected.intensity,
                polarity=selected.polarity,
            ))
            cursor = end
        return matches


def _read_xlsx(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("读取 Excel 需要安装 openpyxl：python -m pip install openpyxl") from error
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = list(workbook.active.iter_rows(values_only=True))
    finally:
        # Windows 会锁定未关闭的 xlsx，导致测试或批处理无法清理文件。
        workbook.close()

    header_index = None
    word_headers = set(HEADER_ALIASES["word"])
    emotion_headers = set(HEADER_ALIASES["emotion"])
    for index, row in enumerate(values[:20]):
        names = {str(value or "").strip() for value in row}
        if names & word_headers and names & emotion_headers:
            header_index = index
            break
    if header_index is None:
        raise ValueError("Excel 前 20 行中未找到词语和情感分类表头")
    headers = [str(value or "").strip() for value in values[header_index]]
    return [
        dict(zip(headers, row))
        for row in values[header_index + 1 :]
        if any(value not in (None, "") for value in row)
    ]


def load_lexicon(path: str | Path) -> EmotionLexicon:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"未找到情感词典：{path}")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        rows = _read_xlsx(path)
    elif path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            rows = list(csv.DictReader(file))
    else:
        raise ValueError("词典仅支持 .csv、.xlsx 或 .xlsm")
    return EmotionLexicon.from_rows(rows)
